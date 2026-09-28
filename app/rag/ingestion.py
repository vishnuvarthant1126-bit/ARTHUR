"""Step 1 of RAG: get clean text out of a file, page by page.

Supported: PDF, DOCX, TXT/MD, CSV. Each file becomes a list of `Page`s so
that answers can later cite *where* a passage came from:

    PDF   -> one Page per real page          -> "handbook.pdf, p. 12"
    CSV   -> one Page per block of rows      -> "grades.csv, rows 2-26"
    DOCX / TXT -> one Page for the whole text -> "notes.docx"  (no page numbers exist)

Everything here is plain parsing: files are never executed, and any parsing
failure becomes a friendly DocumentError instead of a crash.
"""

import csv
import io
import re
from dataclasses import dataclass

SUPPORTED_TYPES = {".pdf": "pdf", ".docx": "docx", ".txt": "text", ".md": "text", ".csv": "csv"}
CSV_ROWS_PER_PAGE = 25


class DocumentError(Exception):
    """A problem with the uploaded file, safe to show the user."""


@dataclass
class Page:
    text: str
    number: int | None = None  # real page number (PDF only)
    label: str | None = None  # how to cite it: "p. 3", "rows 2-26", or None


def extract(data: bytes, extension: str) -> list[Page]:
    """Extract and clean the text of a file. Raises DocumentError on problems."""
    kind = SUPPORTED_TYPES.get(extension.lower())
    if kind is None:
        allowed = ", ".join(sorted(SUPPORTED_TYPES))
        raise DocumentError(f"Unsupported file type '{extension}'. Supported: {allowed}.")
    try:
        match kind:
            case "pdf":
                pages = _pdf(data)
            case "docx":
                pages = _docx(data)
            case "csv":
                pages = _csv(data)
            case _:
                pages = [Page(text=_decode(data))]
    except DocumentError:
        raise
    except Exception as exc:  # corrupt or unusual files: never crash ARTHUR
        raise DocumentError(
            f"Couldn't read this {kind.upper()} file ({type(exc).__name__})."
        ) from exc

    pages = [Page(clean_text(p.text), p.number, p.label) for p in pages]
    pages = [p for p in pages if p.text]
    if not pages:
        hint = (
            " It may be a scanned PDF (images only) - OCR isn't supported yet."
            if kind == "pdf"
            else ""
        )
        raise DocumentError(f"No readable text found in this file.{hint}")
    return pages


def clean_text(text: str) -> str:
    """Tidy extracted text so chunks and embeddings aren't polluted by layout noise."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join words hyphenated across lines
    text = re.sub(r"[ \t\f\v]+", " ", text)  # runs of spaces/tabs -> one space
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)  # keep paragraph breaks, drop big gaps
    # A single line break inside a paragraph is usually just the PDF's line wrapping.
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    return text.strip()


def _pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("This PDF is password-protected. Remove the password and try again.")
    return [
        Page(text=page.extract_text() or "", number=i, label=f"p. {i}")
        for i, page in enumerate(reader.pages, start=1)
    ]


def _docx(data: bytes) -> list[Page]:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page(text="\n\n".join(parts))]


def _csv(data: bytes) -> list[Page]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    pages = []
    for start in range(0, len(body), CSV_ROWS_PER_PAGE):
        block = body[start : start + CSV_ROWS_PER_PAGE]
        # "name: Alice; grade: A" reads far better for an LLM than raw commas.
        lines = [
            "; ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False) if v.strip())
            for row in block
        ]
        first, last = start + 2, start + 1 + len(block)  # +1 header row, 1-based
        pages.append(Page(text="\n\n".join(lines), label=f"rows {first}-{last}"))
    return pages


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")  # never fails
